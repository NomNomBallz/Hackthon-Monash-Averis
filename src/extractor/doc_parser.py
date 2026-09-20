import os
import io
import docx
import pandas as pd
import pdfplumber
import pypdfium2 as pdfium

def parse_attachment(file_path: str) -> dict:
    """
    Parses attachments (.txt, .docx, .xlsx, .pdf).
    Returns text content when available, or image bytes for scanned PDFs.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Attachment file not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return {"type": "text", "content": f.read()}

    elif ext == ".docx":
        doc = docx.Document(file_path)
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return {"type": "text", "content": "\n".join(lines)}

    elif ext in [".xlsx", ".xls"]:
        excel_data = pd.read_excel(file_path, sheet_name=None)
        sheets_text = []
        for sheet_name, df in excel_data.items():
            sheets_text.append(f"--- Sheet: {sheet_name} ---")
            sheets_text.append(df.to_csv(index=False, sep=" "))
        return {"type": "text", "content": "\n".join(sheets_text)}

    elif ext == ".pdf":
        try:
            # 1. Primary path: Try extracting native text using pdfplumber
            text_content = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text(layout=True)
                    if t:
                        text_content.append(t)
            
            extracted_text = "\n".join(text_content).strip()
            
            # If standard text exists, return it
            if len(extracted_text) > 50:
                return {"type": "text", "content": extracted_text}
            
            # 2. Vision Fallback: Render first page into JPEG image bytes using pypdfium2
            pdf = pdfium.PdfDocument(file_path)
            if len(pdf) == 0:
                raise ValueError("PDF has 0 pages.")
                
            page = pdf[0]
            image = page.render(scale=2).to_pil()
            
            img_byte_arr = io.BytesIO()
            image.save(img_byte_arr, format='JPEG')
            
            return {"type": "image", "content": img_byte_arr.getvalue()}

        except Exception as e:
            # Catch corrupt / garbled PDFs (e.g., email_515 missing /Root)
            raise ValueError(f"UNREADABLE_PDF: {str(e)}")

    else:
        raise ValueError(f"Unsupported file format: {ext}")